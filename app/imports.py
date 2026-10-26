"""Importing the Listed Sheet and the Michaels "Download Listing" file."""

from __future__ import annotations

import csv
import io
import json
import re
import secrets
import threading
import time
import traceback
from datetime import datetime

from . import db, rules, titles
from .db import AMAZON_BASE, conn, log_activity, now_str, tx
from .rules import to_bool, to_float

HEADER_KEYS = {
    "ASIN": "asin", "SOURCE LINK": "link", "PRODUCT LINK": "link", "SKU": "sku", "PRODUCT NAME": "name", "CATEGORY": "category",
    "SOURCE PRICE": "source_price", "WEIGHT": "weight", "INVENTORY": "inventory", "LISTED": "listed",
    "LISTING STATUS": "listing_status", "PRICE": "price", "AVAILABILITY": "availability",
    "ESTIMATED DELIVERY": "delivery", "SOURCE": "source", "UPC": "upc", "VARIATION": "variation",
    "STATUS COUNTER": "manual_status", "OPTIMIZER REMARKS": "remarks",
    "EXPORTED MARKETPLACE PRICE": "michaels_price", "EXPORTED MARKETPLACE INVENTORY": "michaels_qty",
    "PRODUCT SOLD 2025": "sold_2025", "PRODUCT SOLD 2026": "sold_2026",
    "SKU NUMBER": "michaels_sku_number",
}


def cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_rows(filename: str, data: bytes):
    if filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise ValueError("Reading .xlsx needs openpyxl. Run setup again, or download the sheet as CSV instead.")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb["Listed Sheet"] if "Listed Sheet" in wb.sheetnames else wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            yield [cell(v) for v in row]
        wb.close()
    else:
        for row in csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))):
            yield row


def map_headers(header: list[str]) -> dict[str, int]:
    idx: dict[str, int] = {}
    for i, h in enumerate(header):
        key = str(h or "").replace("﻿", "").strip().upper()
        if key in HEADER_KEYS and HEADER_KEYS[key] not in idx:
            idx[HEADER_KEYS[key]] = i
        elif key.startswith("LAST OPTIMIZATION RUN") and "run_at" not in idx:
            idx["run_at"] = i
    return idx


def import_products(filename: str, data: bytes, user: str, progress=None) -> dict:
    """progress(stage, done, total), when given, is called as the import moves along (the background import shows it)."""
    note = progress or (lambda *a: None)
    note("Reading the file", 0, 0)
    rows = read_rows(filename, data)
    header = next(rows, None)
    if not header:
        raise ValueError("The file is empty.")
    idx = map_headers(header)
    if "sku" not in idx or "link" not in idx:
        raise ValueError("Couldn't find the SKU and SOURCE LINK columns. Use a download of the Listed Sheet tab.")

    def g(r, k):
        i = idx.get(k)
        return cell(r[i]) if i is not None and i < len(r) else ""

    seen, dup, no_sku, no_link, not_amazon = set(), 0, 0, 0, 0
    records = []
    for n, r in enumerate(rows, start=2):
        sku, link = g(r, "sku"), g(r, "link")
        if not sku:
            if link or g(r, "asin"):
                no_sku += 1
            continue
        if sku in seen:
            dup += 1
            continue
        seen.add(sku)
        asin = rules.asin_from_link(link) or (g(r, "asin") if re.fullmatch(r"[A-Z0-9]{10}", g(r, "asin")) else "")
        if rules.asin_from_link(link):
            link = f"{AMAZON_BASE}/dp/{asin}"
        elif not link:
            no_link += 1
        elif "amazon." not in link.lower():
            not_amazon += 1
        records.append({
            "sku": sku, "asin": asin, "link": link, "name": g(r, "name"), "category": g(r, "category"),
            "source_price": to_float(g(r, "source_price")), "weight": to_float(g(r, "weight")),
            "inventory": to_bool(g(r, "inventory")), "listed": to_bool(g(r, "listed")),
            "listing_status": g(r, "listing_status"), "row_order": n,
            "price": g(r, "price"), "availability": g(r, "availability"), "delivery": g(r, "delivery"),
            "run_at": g(r, "run_at"), "upc": g(r, "upc"), "variation": g(r, "variation"),
            "source": g(r, "source") or "Amazon", "manual_status": g(r, "manual_status"), "remarks": g(r, "remarks"),
            "michaels_price": to_float(g(r, "michaels_price")), "michaels_qty": to_float(g(r, "michaels_qty")),
            "michaels_sku_number": g(r, "michaels_sku_number"),
            "sold_2025": int(to_float(g(r, "sold_2025")) or 0), "sold_2026": int(to_float(g(r, "sold_2026")) or 0),
        })

    note("Saving the products", 0, len(records))
    with tx() as c:
        st = db.get_settings(c)
        pr, tf, tt = db.get_pricing(c), float(st["timeframe_days"]), float(st["title_threshold"])
        existing = {r["sku"]: r for r in c.execute(
            "SELECT sku, price, availability, delivery, run_at, title, title_ok, amazon_variation, created_in_app, ovr FROM products")}
        c.execute("UPDATE products SET active=0 WHERE COALESCE(created_in_app,0)=0")   # app-added products stay active
        added = updated = 0
        for n_done, rec in enumerate(records, start=1):
            if n_done % 400 == 0:
                note("Saving the products", n_done, len(records))
            old = existing.get(rec["sku"])
            # Keep newer scrape data already in the app; otherwise take the sheet's own values.
            if old and old["run_at"] and (not rec["run_at"] or (rules.parse_run(old["run_at"]) or datetime.min)
                                          >= (rules.parse_run(rec["run_at"]) or datetime.min)):
                for k in ("price", "availability", "delivery", "run_at"):
                    rec[k] = old[k]
            rec["fee"], rec["auto_price"] = db.derive_prices(rec["weight"], rec["source_price"],
                                                             db.load_ovr(old["ovr"]) if old else {}, pr)
            rec["price_num"] = rules.scrap_price(rec["price"])
            rec["title_score"], rec["title_note"], rec["title_conflict"] = titles.check(rec["name"], old["title"] if old else "")
            rec.update({"title": old["title"] if old else "", "title_ok": old["title_ok"] if old else "",
                        "amazon_variation": old["amazon_variation"] if old else ""})
            rec["status"] = rules.compute_status(rec, tf, tt)
            if old:
                updated += 1
                c.execute("""UPDATE products SET asin=:asin, link=:link, name=:name, category=:category,
                    source_price=:source_price, weight=:weight, inventory=:inventory, listed=:listed,
                    listing_status=:listing_status, row_order=:row_order, active=1, price=:price,
                    availability=:availability, delivery=:delivery, run_at=:run_at, status=:status,
                    upc=:upc, variation=:variation, source=:source, manual_status=:manual_status, remarks=:remarks,
                    michaels_price=CASE WHEN michaels_updated_at IS NULL THEN :michaels_price ELSE michaels_price END,
                    michaels_qty=CASE WHEN michaels_updated_at IS NULL THEN :michaels_qty ELSE michaels_qty END,
                    michaels_sku_number=CASE WHEN michaels_updated_at IS NULL AND :michaels_sku_number<>''
                        THEN :michaels_sku_number ELSE michaels_sku_number END,
                    sold_2025=:sold_2025, sold_2026=:sold_2026, fee=:fee, auto_price=:auto_price,
                    price_num=:price_num, title_score=:title_score, title_note=:title_note, title_conflict=:title_conflict
                    WHERE sku=:sku""", rec)
            else:
                added += 1
                c.execute("""INSERT INTO products(sku, asin, link, name, category, source_price, weight, inventory,
                    listed, listing_status, row_order, active, price, availability, delivery, run_at, status, upc,
                    variation, source, manual_status, remarks, michaels_price, michaels_qty, sold_2025, sold_2026,
                    created_at, fee, auto_price, price_num, michaels_sku_number)
                    VALUES(:sku, :asin, :link, :name, :category, :source_price, :weight, :inventory, :listed,
                    :listing_status, :row_order, 1, :price, :availability, :delivery, :run_at, :status, :upc,
                    :variation, :source, :manual_status, :remarks, :michaels_price, :michaels_qty, :sold_2025,
                    :sold_2026, :created_at, :fee, :auto_price, :price_num, :michaels_sku_number)""", {**rec, "created_at": now_str()})
        removed = sum(1 for k, r in existing.items() if k not in seen and not r["created_in_app"])
        summary = {"products": len(records), "added": added, "updated": updated, "removed": removed,
                   "duplicate_skus": dup, "rows_without_sku": no_sku, "no_source_link": no_link,
                   "not_amazon_link": not_amazon}
        c.execute("INSERT INTO imports(at, filename, summary) VALUES(?,?,?)", (now_str(), filename, json.dumps(summary)))
        log_activity(c, user, "imported the Listed Sheet", "", f"{len(records):,} products from {filename}")
    note("Updating the word weights", len(records), len(records))
    db.rebuild_title_model()           # the catalog's words changed: relearn which ones identify a product
    return summary


# ------------------------------------------------------------------ the import as a background job
# A 32,000-row import takes a while on a small server. Run through a normal request, a proxy or the browser can give up
# before the answer arrives, and the screen is left waiting although the import finished. As a job, the browser only
# starts it and then asks how it is going; the result is kept on the server, so it can be read after a reload or a drop.
JOB_KEY = "import_job"
_JOB: dict = {}
_JOB_LOCK = threading.Lock()


def _save_job(job: dict) -> None:
    with tx() as c:
        db.set_setting(c, JOB_KEY, json.dumps(job))


def job_status() -> dict:
    with _JOB_LOCK:
        if _JOB:
            return dict(_JOB)
    row = conn().execute("SELECT value FROM settings WHERE key=?", (JOB_KEY,)).fetchone()
    try:
        return json.loads(row["value"]) if row else {"state": "idle"}
    except ValueError:
        return {"state": "idle"}


def job_reset() -> None:
    """At startup: a job saved as running was cut off by the restart."""
    job = job_status()
    if job.get("state") == "running":
        _save_job({**job, "state": "error", "finished_at": now_str(), "error": "The server restarted while this import was running. Nothing was changed; import the file again."})
    with _JOB_LOCK:
        _JOB.clear()


def _run_job(filename: str, data: bytes, user: str) -> None:
    def progress(stage, done, total):
        with _JOB_LOCK:
            _JOB.update({"stage": stage, "done": done, "total": total})
    try:
        summary = import_products(filename, data, user, progress)
        end = {"state": "done", "summary": summary}
    except ValueError as e:
        end = {"state": "error", "error": str(e)}
    except Exception as e:                                       # keep the server log's full traceback, and tell the screen what kind of problem it was
        traceback.print_exc()
        end = {"state": "error", "error": f"The import stopped because of an unexpected problem ({type(e).__name__}: {e}). The server log has the details."}
    with _JOB_LOCK:
        _JOB.update({**end, "finished_at": now_str()})
        final = dict(_JOB)
    _save_job(final)


def start_job(filename: str, data: bytes, user: str) -> dict:
    """Start importing in the background and return at once. Only one import can run at a time."""
    if not data:
        raise ValueError("The file is empty.")
    with _JOB_LOCK:
        if _JOB.get("state") == "running":
            raise ValueError("An import is already running. Wait for it to finish.")
        _JOB.clear()
        _JOB.update({"state": "running", "filename": filename, "user": user, "started_at": now_str(),
                     "stage": "Reading the file", "done": 0, "total": 0})
        first = dict(_JOB)
    _save_job(first)
    threading.Thread(target=_run_job, args=(filename, data, user), daemon=True, name="import").start()
    return first


# ------------------------------------------------------------------ Michaels "Download Listing"
MZ_FIELDS = {
    "sku": ["sku", "sellersku", "merchantsku", "vendorsku", "partnersku", "productsku", "offersku", "yoursku", "shopsku"],
    "sku_number": ["skunumber", "michaelsskunumber", "michaelssku", "skuno", "skunum", "itemnumber", "michaelsitemnumber",
                   "productid", "listingid", "itemid", "skuid"],
    "price": ["price", "sellingprice", "saleprice", "listprice", "yourprice", "offerprice", "unitprice", "currentprice",
              "retailprice"],
    "qty": ["qty", "quantity", "inventory", "stock", "availablequantity", "quantityavailable", "availableqty",
            "onhand", "stockquantity", "inventoryquantity"],
}
MZ_PENDING: dict[str, dict] = {}   # token -> parsed upload waiting for "Apply"

# Product image columns: any header naming an image/photo/picture/thumbnail (one column per image, or one column
# holding several links). Links are split on spaces, commas, semicolons and pipes.
IMG_WORDS = ("image", "img", "photo", "picture", "thumbnail")
IMG_SKIP = ("count", "alt", "status", "valid", "error", "note", "type", "size", "width", "height", "quality", "required")
URL_SPLIT = re.compile(r"[\s,;|]+")


def parse_images(row, idxs) -> list[str]:
    urls: list[str] = []
    for i in idxs:
        if i < len(row):
            for tok in URL_SPLIT.split(str(row[i]).strip()):
                if tok.startswith("//"):
                    tok = "https:" + tok
                if re.match(r"https?://", tok, re.I) and tok not in urls:
                    urls.append(tok)
    return urls[:20]


def _norm(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def mz_detect(header: list[str]) -> dict[str, int]:
    norm = [_norm(h) for h in header]
    found: dict[str, int] = {}
    for field, names in MZ_FIELDS.items():
        for n in names:                                   # exact names first
            if n in norm and norm.index(n) not in found.values():
                found[field] = norm.index(n)
                break
        if field not in found:                            # then "contains"
            for i, h in enumerate(norm):
                if i in found.values() or (field == "sku" and "number" in h):
                    continue
                if any(n in h for n in names):
                    found[field] = i
                    break
    used = set(v for v in found.values() if isinstance(v, int))
    imgs = [i for i, h in enumerate(norm) if i not in used and any(w in h for w in IMG_WORDS)
            and not any(s in h for s in IMG_SKIP)]
    if imgs:
        found["images"] = imgs
    return found


def mz_read(filename: str, data: bytes):
    if filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise ValueError("Reading .xlsx needs openpyxl. Run setup again, or upload the CSV version.")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [["" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() and abs(v) < 1e15 else str(v))
                 for v in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
    else:
        text = data.decode("utf-8-sig", errors="replace")
        delim = "\t" if text[:5000].count("\t") > text[:5000].count(",") else ","
        rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ValueError("The file is empty.")
    hi = next((i for i, r in enumerate(rows[:15]) if any("sku" in _norm(c) for c in r)), 0)   # first row naming a SKU column
    return rows[hi], rows[hi + 1:]


def mz_preview(filename: str, data: bytes, mapping: dict | None = None) -> dict:
    header, rows = mz_read(filename, data)
    cols = mz_detect(header)
    if mapping:
        cols = {k: int(v) for k, v in mapping.items() if k != "images" and str(v).lstrip("-").isdigit() and int(v) >= 0}
        imgs = mapping.get("images")
        imgs = imgs if isinstance(imgs, list) else [x for x in str(imgs or "").split(",") if x.strip()]
        imgs = [int(x) for x in imgs if str(x).strip().lstrip("-").isdigit() and int(x) >= 0]
        if imgs:
            cols["images"] = imgs
    img_cols = cols.get("images") or []
    missing = [k for k in ("sku", "price", "qty") if k not in cols]
    base = {"headers": header, "mapping": cols, "filename": filename}
    if missing:
        labels = {"sku": "SKU", "price": "Price", "qty": "Quantity"}
        return {**base, "ok": False, "error": "Couldn't find the " + ", ".join(labels[m] for m in missing) + " column. Pick it below."}

    def g(r, k):
        i = cols.get(k)
        return str(r[i]).strip() if i is not None and i < len(r) else ""

    parsed, dup, blank, rounded, rounded_skus = {}, 0, 0, 0, []
    for r in rows:
        sku = g(r, "sku")
        if not sku:
            blank += 1
            continue
        key = sku.lower()
        if key in parsed:
            dup += 1
            continue
        num = g(r, "sku_number")
        if re.search(r"[eE]\+?\d+$", num) or re.fullmatch(r"\d+\.0+", num):
            rounded += 1          # Excel turned a long number into 4.36E+17 and lost the digits: never store it
            rounded_skus.append(key)
            num = ""              # blank = keep whatever number the catalog already has
        qty = to_float(g(r, "qty"))
        parsed[key] = {"sku": sku, "sku_number": num, "price": to_float(g(r, "price")),
                       "qty": int(qty) if qty is not None else None, "images": parse_images(r, img_cols)}
    prods = {r["sku"].lower(): r for r in conn().execute(
        "SELECT sku, name, listed, michaels_sku_number, michaels_price, michaels_qty FROM products WHERE active=1")}
    matched = [k for k in parsed if k in prods]
    changes, price_changes, qty_changes, new_links = [], 0, 0, 0
    for k in matched:
        p, m = prods[k], parsed[k]
        pc = m["price"] is not None and (p["michaels_price"] is None or abs((p["michaels_price"] or 0) - m["price"]) > 0.001)
        qc = m["qty"] is not None and (p["michaels_qty"] is None or int(p["michaels_qty"]) != m["qty"])
        nc = bool(m["sku_number"]) and (p["michaels_sku_number"] or "") != m["sku_number"]
        price_changes += pc
        qty_changes += qc
        new_links += nc
        if (pc or qc) and len(changes) < 60:
            changes.append({"sku": p["sku"], "name": p["name"], "old_price": p["michaels_price"], "new_price": m["price"],
                            "old_qty": p["michaels_qty"], "new_qty": m["qty"]})
    # rounded numbers are ignored; how many of those products already have a good number, and how many don't?
    rounded_matched = [k for k in rounded_skus if k in prods]
    rounded_kept = sum(1 for k in rounded_matched if prods[k]["michaels_sku_number"])
    not_in_catalog = [parsed[k]["sku"] for k in parsed if k not in prods]
    missing_listed = [p["sku"] for k, p in prods.items() if k not in parsed and p["listed"]]
    will_clear = conn().execute("""SELECT COUNT(*) FROM products WHERE active=1 AND (michaels_price IS NOT NULL
        OR michaels_qty IS NOT NULL OR michaels_updated_at IS NOT NULL)""").fetchone()[0] - sum(
        1 for k in matched if prods[k]["michaels_price"] is not None or prods[k]["michaels_qty"] is not None)
    token = secrets.token_hex(12)
    MZ_PENDING.clear()   # only the latest preview is kept
    MZ_PENDING[token] = {"filename": filename, "rows": parsed, "at": time.time(), "matched": len(matched),
                         "has_images": bool(img_cols)}
    return {**base, "ok": True, "token": token, "rows": len(parsed), "matched": len(matched),
            "will_clear": max(0, will_clear), "has_images": bool(img_cols),
            "image_columns": [header[i] for i in img_cols if i < len(header)],
            "with_images": sum(1 for k in matched if parsed[k]["images"]),
            "image_urls": sum(len(parsed[k]["images"]) for k in matched),
            "price_changes": price_changes, "qty_changes": qty_changes, "sku_numbers_new": new_links,
            "duplicates": dup, "blank_sku": blank, "rounded_sku_numbers": rounded,
            "rounded_kept": rounded_kept, "rounded_missing": len(rounded_matched) - rounded_kept,
            "not_in_catalog": len(not_in_catalog), "not_in_catalog_sample": not_in_catalog[:40],
            "missing_listed": len(missing_listed), "missing_listed_sample": missing_listed[:40],
            "has_sku_number": "sku_number" in cols, "changes": changes}


def mz_apply(token: str, clear_missing: bool, user: str) -> dict:
    pending = MZ_PENDING.get(token or "")
    if not pending or time.time() - pending["at"] > 3600:
        raise ValueError("This preview has expired. Upload the file again.")
    if not pending["matched"]:
        raise ValueError("None of the rows in this file match a product SKU, so nothing was changed. "
                         "Check that you picked the right file and the right SKU column.")
    rows, now = pending["rows"], now_str()
    has_images = pending.get("has_images", False)
    updated = cleared = with_images = 0
    with tx() as c:
        for p in c.execute("SELECT sku FROM products WHERE active=1").fetchall():
            m = rows.get(p["sku"].lower())
            if m:
                c.execute("""UPDATE products SET michaels_sku_number=COALESCE(NULLIF(?,''), michaels_sku_number),
                    michaels_price=?, michaels_qty=?, michaels_updated_at=? WHERE sku=?""",
                          (m["sku_number"], m["price"], m["qty"], now, p["sku"]))
                if has_images:         # the images always follow the latest download, like price and quantity
                    c.execute("UPDATE products SET images=? WHERE sku=?",
                              (json.dumps(m["images"]) if m["images"] else None, p["sku"]))
                    with_images += bool(m["images"])
                updated += 1
            elif clear_missing:
                cleared += c.execute("""UPDATE products SET michaels_price=NULL, michaels_qty=NULL, michaels_updated_at=NULL
                    WHERE sku=? AND (michaels_price IS NOT NULL OR michaels_qty IS NOT NULL
                    OR michaels_updated_at IS NOT NULL)""", (p["sku"],)).rowcount
                if has_images:
                    c.execute("UPDATE products SET images=NULL WHERE sku=?", (p["sku"],))
        summary = {"rows": len(rows), "updated": updated, "cleared": cleared, "clear_missing": clear_missing,
                   "with_images": with_images}
        c.execute("INSERT INTO michaels_imports(at, user, filename, summary) VALUES(?,?,?,?)",
                  (now, user, pending["filename"], json.dumps(summary)))
        log_activity(c, user, "imported the Michaels download", "",
                     f"{updated:,} products updated from {pending['filename']}" +
                     (f", {cleared:,} not in the file were cleared" if cleared else ""))
    MZ_PENDING.clear()
    return {"ok": True, **summary}


def mz_history() -> dict:
    return {"items": [{**dict(r), "summary": json.loads(r["summary"])} for r in
                      conn().execute("SELECT * FROM michaels_imports ORDER BY id DESC LIMIT 20")]}
