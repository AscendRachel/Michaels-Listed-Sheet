"""Sold counts: the app's version of the sheet's IMPORTRANGE + COUNTIF formulas.

In the Google Sheet, PRODUCT SOLD 2025 / 2026 are `COUNTIF(IMPORTRANGE(tracker, "Product sold tracker!A4:L"), "=" & SKU)`:
the number of cells in the tracker's range that equal the product's SKU (case ignored). That is what is done here.
Each year has a link to its tracker sheet and a range. The sheet must be shared as "Anyone with the link can view" so
the server can read it without a Google login; otherwise download the tab as CSV/XLSX and upload it instead.
TOTAL PRODUCT SOLD stays 2025 + 2026 (it is calculated, not stored).
"""

from __future__ import annotations

import csv
import io
import json
import re
import threading
import time
import urllib.error
import urllib.request
from collections import Counter
from datetime import datetime

from . import db
from .db import conn, log_activity, now_str, tx

YEARS = ("2025", "2026")
DEFAULT_RANGES = {"2025": "A4:L", "2026": "A3:H"}
CONFIG_KEY, STATUS_KEY = "sold_sources", "sold_status"
SHEET_URL = re.compile(r"^https://docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]+)")
GID = re.compile(r"[#&?]gid=(\d+)")
RANGE = re.compile(r"^\s*([A-Za-z]{1,3})(\d*)\s*:\s*([A-Za-z]{1,3})(\d*)\s*$")
MAX_BYTES = 60 * 1024 * 1024
TAB_NAME = "Product sold tracker"


# ------------------------------------------------------------------ settings
def get_config(c=None) -> dict:
    c = c or conn()
    row = c.execute("SELECT value FROM settings WHERE key=?", (CONFIG_KEY,)).fetchone()
    try:
        saved = json.loads(row["value"]) if row else {}
    except ValueError:
        saved = {}
    cfg = {"auto": bool(saved.get("auto", False)), "every_hours": int(saved.get("every_hours") or 6)}
    for y in YEARS:
        s = saved.get(y) or {}
        cfg[y] = {"url": str(s.get("url") or ""), "range": str(s.get("range") or DEFAULT_RANGES[y])}
    return cfg


def get_status(c=None) -> dict:
    c = c or conn()
    row = c.execute("SELECT value FROM settings WHERE key=?", (STATUS_KEY,)).fetchone()
    try:
        return json.loads(row["value"]) if row else {}
    except ValueError:
        return {}


def save_config(b: dict, user: str) -> dict:
    cfg = get_config()
    for y in YEARS:
        s = b.get(y) or {}
        url = str(s.get("url") or "").strip()
        if url and not SHEET_URL.match(url):
            raise ValueError(f"The {y} link must be a Google Sheets link (https://docs.google.com/spreadsheets/d/…).")
        rng = str(s.get("range") or DEFAULT_RANGES[y]).strip().upper()
        parse_range(rng)
        cfg[y] = {"url": url, "range": rng}
    if "auto" in b:
        cfg["auto"] = bool(b["auto"])
    if b.get("every_hours"):
        cfg["every_hours"] = max(1, min(168, int(b["every_hours"])))
    with tx() as c:
        db.set_setting(c, CONFIG_KEY, json.dumps(cfg))
        log_activity(c, user, "changed the sold-count sources", "", ", ".join(y for y in YEARS if cfg[y]["url"]))
    return cfg


# ------------------------------------------------------------------ reading a tracker
def col_index(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + ord(ch) - 64
    return n - 1


def parse_range(text: str) -> tuple[int, int, int, int | None]:
    """'A4:L' -> (first column, last column, first row, last row or None). Rows count from 1."""
    m = RANGE.match(text or "")
    if not m:
        raise ValueError("A range looks like A4:L (first cell, then the last column).")
    c1, c2 = col_index(m.group(1)), col_index(m.group(3))
    if c2 < c1:
        c1, c2 = c2, c1
    return c1, c2, int(m.group(2) or 1), int(m.group(4)) if m.group(4) else None


def count_values(rows: list[list[str]], rng: str) -> tuple[Counter, int]:
    """Like COUNTIF over a range: how many cells equal each value (case ignored). Also the number of non-empty rows read."""
    c1, c2, r1, r2 = parse_range(rng)
    counts: Counter = Counter()
    used = 0
    for row in rows[r1 - 1:r2]:
        seen = False
        for v in row[c1:c2 + 1]:
            v = str(v).strip().lower()
            if v:
                counts[v] += 1
                seen = True
        used += seen
    return counts, used


def export_url(url: str) -> str:
    m = SHEET_URL.match(url or "")
    if not m:
        raise ValueError("That isn't a Google Sheets link.")
    g = GID.search(url)
    return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv" + (f"&gid={g.group(1)}" if g else "")


def fetch_rows(url: str) -> list[list[str]]:
    """The tab behind a Google Sheets link, as rows of text. Needs 'Anyone with the link can view'."""
    req = urllib.request.Request(export_url(url), headers={"User-Agent": "Mozilla/5.0 (Greentek Catalog sold-count sync)"})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            ctype = r.headers.get("Content-Type", "")
            body = r.read(MAX_BYTES + 1)
    except urllib.error.HTTPError as e:
        raise ValueError("Google refused the request (" + str(e.code) + "). Is the link right, and is the sheet shared as 'Anyone with the link can view'?")
    except (urllib.error.URLError, TimeoutError, OSError) as e:
        raise ValueError(f"Couldn't reach Google Sheets: {getattr(e, 'reason', e)}")
    if len(body) > MAX_BYTES:
        raise ValueError("That sheet is too large to read.")
    if "csv" not in ctype.lower():
        raise ValueError("Google asked for a sign-in, so the sheet isn't shared by link. Open it, Share → General access → "
                         "'Anyone with the link' → Viewer, or download the tab and upload the file instead.")
    return list(csv.reader(io.StringIO(body.decode("utf-8-sig", errors="replace"))))


def read_file(filename: str, data: bytes) -> list[list[str]]:
    from .imports import cell
    if filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise ValueError("Reading .xlsx needs openpyxl. Download the tab as CSV instead.")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[TAB_NAME] if TAB_NAME in wb.sheetnames else wb.worksheets[0]
        rows = [[cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    return list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))


# ------------------------------------------------------------------ syncing
def sync(user: str, years: list[str] | None = None, files: dict | None = None) -> dict:
    """Read each year's tracker and set PRODUCT SOLD <year> for every product. Nothing is changed if a tracker
    can't be read, is empty, or contains none of the catalog's SKUs."""
    cfg, files = get_config(), files or {}
    todo = [y for y in (years or YEARS) if y in YEARS and (y in files or cfg[y]["url"])]
    if not todo:
        raise ValueError("Add the link to a tracker sheet first (or upload its file).")
    results: dict = {}
    pulled: dict[str, tuple[Counter, int, str]] = {}
    for y in todo:
        try:
            if y in files:
                rows, source = read_file(*files[y]), files[y][0]
            else:
                rows, source = fetch_rows(cfg[y]["url"]), "Google Sheets link"
            counts, used = count_values(rows, cfg[y]["range"])
            if not used:
                raise ValueError(f"Nothing found in {cfg[y]['range']} of the {y} sheet. Is it the right tab and range?")
            pulled[y] = (counts, used, source)
        except ValueError as e:
            results[y] = {"error": str(e)}
    with tx() as c:
        for y, (counts, used, source) in pulled.items():
            col = f"sold_{y}"
            rows = c.execute(f"SELECT sku, {col} AS old FROM products").fetchall()
            new = {r["sku"]: counts.get(r["sku"].strip().lower(), 0) for r in rows}
            matched = sum(1 for v in new.values() if v)
            had = sum(1 for r in rows if (r["old"] or 0) > 0)
            if not matched and had:
                results[y] = {"error": f"None of your SKUs were found in the {y} sheet, so nothing was changed. Wrong tab or range?"}
                continue
            old = {r["sku"]: r["old"] or 0 for r in rows}
            changes = [(v, sku) for sku, v in new.items() if v != old[sku]]
            c.executemany(f"UPDATE products SET {col}=? WHERE sku=?", changes)
            results[y] = {"rows": used, "matched": matched, "sales": sum(new.values()), "changed": len(changes), "source": source}
        ok = {y: r for y, r in results.items() if "error" not in r}
        status = {"at": now_str(), "by": user, "years": results, "ok": bool(ok) and len(ok) == len(results)}
        db.set_setting(c, STATUS_KEY, json.dumps(status))
        if ok:
            log_activity(c, user, "synced the sold counts", "", "; ".join(f"{y}: {r['sales']} sold, {r['changed']} products changed" for y, r in ok.items()))
    return status


def auto_loop() -> None:
    """Sync on a schedule when it is switched on (checked every 10 minutes; the PC has to be on)."""
    while True:
        time.sleep(600)
        try:
            cfg = get_config()
            if not cfg["auto"] or not any(cfg[y]["url"] for y in YEARS):
                continue
            last = get_status().get("at")
            if last:
                try:
                    if (datetime.now() - datetime.strptime(last, "%Y-%m-%d %H:%M:%S")).total_seconds() < cfg["every_hours"] * 3600:
                        continue
                except ValueError:
                    pass
            sync("Auto sync")
        except Exception as e:                                          # never let a failed sync kill the thread
            print(f"Sold-count sync failed: {e}", flush=True)


def start_background() -> None:
    threading.Thread(target=auto_loop, daemon=True).start()
